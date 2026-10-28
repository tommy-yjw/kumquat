#!/usr/bin/env python3
"""pdf2docx 输出修复:合并被分页切碎的表格 + 按字号重建标题层级。
用法: python3 pdf2docx_postfix.py <file.docx>"""
import sys, copy
from docx import Document

def is_tbl(el): return el.tag.endswith('}tbl')
def is_p(el): return el.tag.endswith('}p')

def merge_tables(doc):
    body = doc.element.body
    by_el = {t._tbl: t for t in doc.tables}
    merged = 0
    prev = None          # 上一个"可续接"的表格
    gap = 0              # 与上一个表格之间隔的空段落
    for el in list(body):
        if is_tbl(el):
            cur = by_el.get(el)
            if cur is None:
                prev, gap = None, 0
                continue
            if (prev is not None and gap <= 1
                    and len(cur.columns) == len(prev.columns)):
                first_texts = [c.text.strip() for c in prev.rows[0].cells]
                for row in list(cur.rows):
                    texts = [c.text.strip() for c in row.cells]
                    if texts == first_texts:   # 跨页重复表头:跳过
                        continue
                    prev._tbl.append(copy.deepcopy(row._tr))
                body.remove(el)
                merged += 1
                gap = 0                        # prev 保持,继续吞并下一块
            else:
                prev, gap = cur, 0
        elif is_p(el):
            if prev is not None and ''.join(el.itertext()).strip() == '':
                gap += 1
                if gap > 1:
                    prev, gap = None, 0
            else:
                prev, gap = None, 0
        else:
            prev, gap = None, 0
    return merged

def rebuild_headings(doc):
    count = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text or len(text) > 80:
            continue
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if not sizes:
            continue
        size = max(sizes)
        level = None
        if size >= 17:
            level = 1
        elif size >= 14:
            level = 2
        elif size >= 12.5:
            level = 3
        if level:
            try:
                p.style = doc.styles['Heading %d' % level]
                count += 1
            except KeyError:
                break
    return count

def main():
    path = sys.argv[1]
    doc = Document(path)
    merged = merge_tables(doc)
    headings = rebuild_headings(doc)
    doc.save(path)
    print(f"postfix: 合并表格碎片 {merged} 处,重建标题 {headings} 个")

if __name__ == '__main__':
    main()
